"""
Phase F4: verify /api/finance/export/{year|month}.{pdf|xlsx} endpoints.

Cover:
- Year PDF returns valid %PDF- bytes + Content-Disposition inline
- Year XLSX returns valid ZIP (openxml) bytes + Content-Disposition attachment
- Month PDF + Month XLSX same shape
- Auth via ?token= query (authenticate_file_token)
- Role gating: sales -> 403; admin/finance/management -> 200
- Year clamp: year=1999 -> fallback ke current_year (tetap 200, tidak error)
"""
from tests.conftest import bearer


def _assert_pdf(resp):
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"].startswith("application/pdf")
    body = resp.content
    assert body[:5] == b"%PDF-", f"bukan file PDF: {body[:20]!r}"
    assert len(body) > 500
    cd = resp.headers.get("content-disposition", "")
    assert cd.startswith("inline"), cd


def _assert_xlsx(resp):
    assert resp.status_code == 200, resp.text
    assert "spreadsheetml.sheet" in resp.headers["content-type"]
    body = resp.content
    # XLSX = ZIP archive => magic PK\x03\x04
    assert body[:4] == b"PK\x03\x04", f"bukan XLSX (zip): {body[:8]!r}"
    assert len(body) > 500
    cd = resp.headers.get("content-disposition", "")
    assert cd.startswith("attachment"), cd


def test_export_year_pdf_via_query_token(client, admin_token):
    r = client.get(f"/api/finance/export/year.pdf?year=2026&token={admin_token}")
    _assert_pdf(r)


def test_export_year_xlsx_via_query_token(client, admin_token):
    r = client.get(f"/api/finance/export/year.xlsx?year=2026&token={admin_token}")
    _assert_xlsx(r)


def test_export_month_pdf_via_query_token(client, admin_token):
    r = client.get(
        f"/api/finance/export/month.pdf?year=2026&month=9&token={admin_token}")
    _assert_pdf(r)


def test_export_month_xlsx_via_query_token(client, admin_token):
    r = client.get(
        f"/api/finance/export/month.xlsx?year=2026&month=9&token={admin_token}")
    _assert_xlsx(r)


def test_export_year_pdf_via_bearer_header(client, admin_token):
    """authenticate_file_token juga terima header Authorization biasa."""
    r = client.get("/api/finance/export/year.pdf?year=2026",
                   headers=bearer(admin_token))
    _assert_pdf(r)


def test_export_year_pdf_denied_without_token(client):
    r = client.get("/api/finance/export/year.pdf?year=2026")
    assert r.status_code == 401


def test_export_year_pdf_denied_for_sales(client, sales_token):
    """Sales tidak boleh export -- endpoint restricted ke admin/finance/management."""
    r = client.get(
        f"/api/finance/export/year.pdf?year=2026&token={sales_token}")
    assert r.status_code == 403


def test_export_year_xlsx_denied_for_sales(client, sales_token):
    r = client.get(
        f"/api/finance/export/year.xlsx?year=2026&token={sales_token}")
    assert r.status_code == 403


def test_export_year_clamps_out_of_range(client, admin_token):
    """Phase 14a-2 style clamp: year<2000 fallback ke current_year, tidak error."""
    r = client.get(
        f"/api/finance/export/year.pdf?year=1999&token={admin_token}")
    _assert_pdf(r)


def test_export_month_clamps_out_of_range(client, admin_token):
    """month<1 atau >12 fallback ke current month, tidak error."""
    r = client.get(
        f"/api/finance/export/month.pdf?year=2026&month=99&token={admin_token}")
    _assert_pdf(r)


# ============================================================================
# Phase EX-6: Project Breakdown section muncul di PDF + XLSX export.
# ============================================================================


def test_ex6_xlsx_month_has_project_breakdown_sheet(client, admin_token):
    r = client.get(
        f"/api/finance/export/month.xlsx?year=2026&month=9&token={admin_token}")
    _assert_xlsx(r)
    import io
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(r.content), read_only=True)
    assert "Project Breakdown" in wb.sheetnames, \
        f"Sheet 'Project Breakdown' tidak ada. Ada: {wb.sheetnames}"


def test_ex6_xlsx_year_has_project_breakdown_sheet(client, admin_token):
    r = client.get(
        f"/api/finance/export/year.xlsx?year=2026&token={admin_token}")
    _assert_xlsx(r)
    import io
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(r.content), read_only=True)
    assert "Project Breakdown" in wb.sheetnames


def test_ex6_pdf_month_still_valid(client, admin_token):
    """Regression: PDF Bulan valid setelah tambahan section Project."""
    r = client.get(
        f"/api/finance/export/month.pdf?year=2026&month=9&token={admin_token}")
    _assert_pdf(r)


def test_ex6_pdf_year_still_valid(client, admin_token):
    r = client.get(
        f"/api/finance/export/year.pdf?year=2026&token={admin_token}")
    _assert_pdf(r)
