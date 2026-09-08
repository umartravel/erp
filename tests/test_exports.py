"""
Phase 8a-5: Excel export endpoints -- pastikan file valid, RBAC ketat, dan
content mengandung header + row data yang diharapkan.

Verifikasi:
- Response 200 + Content-Type xlsx + Content-Disposition attachment
- File dapat dibuka openpyxl + sheet punya header row
- RBAC: sales tidak boleh (403)
- Validasi month=YYYY-MM (400 kalau invalid)
- package-jamaah 404 kalau id tidak ada
"""
from io import BytesIO

from openpyxl import load_workbook

from tests.conftest import bearer

import db


_XLSX_MT = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _open_xlsx(response):
    """Load bytes response ke Workbook untuk assertion."""
    return load_workbook(BytesIO(response.content), read_only=True)


def _make_jamaah(client, admin_token, nik, name, order_date):
    """Register jamaah dgn order_date custom (via direct DB update after create)."""
    pkg = client.get("/api/packages", headers=bearer(admin_token)).json()[0]
    r = client.post("/api/jamaah", json={
        "nik": nik, "name": name, "phone": "0812" + nik[-8:],
        "package_type": pkg["name"],
        "total_price": pkg.get("price_quad") or pkg.get("price") or 25000000,
        "status": "Terdaftar",
    }, headers=bearer(admin_token))
    assert r.status_code == 200, r.text
    jid = r.json()["id"]
    # Set order_date manual biar tes bulan tertentu deterministik.
    db.execute("UPDATE jamaah SET order_date = ? WHERE id = ?", (order_date, jid))
    return jid


def test_jamaah_monthly_export_ok(client, admin_token):
    """Jamaah di Jan 2025 tampil di sheet."""
    _make_jamaah(client, admin_token, "9999999999820001", "TEST Export Jan", "2025-01-15")
    _make_jamaah(client, admin_token, "9999999999820002", "TEST Export Feb", "2025-02-01")

    r = client.get("/api/exports/jamaah-monthly.xlsx?month=2025-01",
                   headers=bearer(admin_token))
    assert r.status_code == 200
    assert r.headers["content-type"].startswith(_XLSX_MT)
    assert "closingan-jamaah-2025-01.xlsx" in r.headers.get("content-disposition", "")

    wb = _open_xlsx(r)
    ws = wb.worksheets[0]
    # Row 1 = header
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    assert "NIK" in headers and "Nama" in headers and "Sales" in headers
    # Data rows -- cari jamaah TEST Jan yg dibuat.
    all_names = [c.value for row in ws.iter_rows(min_row=2) for c in row]
    assert "TEST Export Jan" in all_names
    assert "TEST Export Feb" not in all_names  # feb tidak muncul


def test_jamaah_monthly_rbac(client, sales_token):
    """Sales tidak boleh export."""
    r = client.get("/api/exports/jamaah-monthly.xlsx?month=2025-01",
                   headers=bearer(sales_token))
    assert r.status_code == 403


def test_month_invalid_format(client, admin_token):
    """Parameter month bukan YYYY-MM -> 400."""
    r = client.get("/api/exports/jamaah-monthly.xlsx?month=januari-2025",
                   headers=bearer(admin_token))
    assert r.status_code == 400
    assert "YYYY-MM" in r.json()["error"]


def test_packages_monthly_export(client, admin_token, ops_token):
    """Ops boleh export paket bulanan (bukan cuma finance)."""
    r = client.get("/api/exports/packages-monthly.xlsx?month=2026-08",
                   headers=bearer(ops_token))
    assert r.status_code == 200
    wb = _open_xlsx(r)
    # 2 sheet: Berangkat + Pulang
    assert len(wb.worksheets) == 2


def test_finance_monthly_export(client, admin_token, finance_token):
    """Finance boleh export -- 3 sheet (Ringkasan + Pemasukan + Pengeluaran)."""
    r = client.get("/api/exports/finance-monthly.xlsx?month=2026-08",
                   headers=bearer(finance_token))
    assert r.status_code == 200
    wb = _open_xlsx(r)
    assert len(wb.worksheets) == 3
    # Sheet pertama = Ringkasan dgn 3 row metrik (Total Pemasukan, Pengeluaran, Selisih)
    ws = wb.worksheets[0]
    metrics = [c.value for row in ws.iter_rows(min_row=2, min_col=1, max_col=1) for c in row]
    assert "Total Pemasukan" in metrics
    assert "Total Pengeluaran" in metrics
    assert "Selisih (Net)" in metrics


def test_package_jamaah_export(client, admin_token):
    """Klik paket -> jamaah-nya keluar."""
    pkg = client.get("/api/packages", headers=bearer(admin_token)).json()[0]
    r = client.get(f"/api/exports/package-jamaah.xlsx?package_id={pkg['id']}",
                   headers=bearer(admin_token))
    assert r.status_code == 200
    wb = _open_xlsx(r)
    ws = wb.worksheets[0]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    assert "NIK" in headers and "Nama" in headers and "No. Paspor" in headers


def test_package_jamaah_404(client, admin_token):
    r = client.get("/api/exports/package-jamaah.xlsx?package_id=999999",
                   headers=bearer(admin_token))
    assert r.status_code == 404
